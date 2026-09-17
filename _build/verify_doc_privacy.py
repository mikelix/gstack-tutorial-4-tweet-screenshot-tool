# -*- coding: utf-8 -*-
"""Privacy scan on the BUILT docx files, not source content.

Mirrors verify_deck.py's discipline for the pptx deck: a clean source
document doesn't guarantee a clean built artifact. Opens the actual
gstack-tutorial-4_{EN,ZH}.docx files via python-docx and scans every
paragraph and table-cell run for the same forbidden identifiers the source
privacy review flagged (docs/TUTORIAL_4_PRIVACY_REVIEW.md).

Exit 0 if clean, 1 if any forbidden identifier is found.
"""
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"

FORBIDDEN = [
    "meyklee",
    "MEE2159",
    "hkustmech",
    "RTXG-KKGR",
    "oauth/device",
    "user_code",
    r"C:\Users",
]

failures = []


def all_text(doc: Document):
    for p in doc.paragraphs:
        yield p.text
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p.text


def scan(path: Path):
    print(f"\n{path.name}")
    if not path.exists():
        print("  FAIL  file does not exist")
        failures.append(f"{path.name}: missing")
        return
    doc = Document(str(path))
    full_text = "\n".join(all_text(doc))
    hit_any = False
    for pattern in FORBIDDEN:
        if pattern.lower() in full_text.lower():
            print(f"  FAIL  forbidden identifier present: {pattern!r}")
            failures.append(f"{path.name}: {pattern}")
            hit_any = True
    if not hit_any:
        print(f"  OK    zero forbidden identifiers ({len(full_text)} chars scanned)")

    section_count = sum(1 for p in doc.paragraphs if p.text and
                         p.runs and p.runs[0].font.size and
                         p.runs[0].font.size.pt == 14)
    print(f"        {section_count} H1-styled headings, {len(doc.tables)} tables")


if __name__ == "__main__":
    scan(DIST / "gstack-tutorial-4_EN.docx")
    scan(DIST / "gstack-tutorial-4_ZH.docx")

    print("\n" + "=" * 52)
    if failures:
        print(f"FAILED — {len(failures)} issue(s) found.")
        sys.exit(1)
    print("PASSED — both documents clean, zero forbidden identifiers.")
    sys.exit(0)
