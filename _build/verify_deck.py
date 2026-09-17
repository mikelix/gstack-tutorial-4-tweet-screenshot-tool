# -*- coding: utf-8 -*-
"""Verify the built decks: slide count, EN/ZH parity, and a privacy scan of
the ACTUAL RENDERED TEXT (not the source content file) -- checking the
generator's output, not just trusting deck_content.py was clean, is the
same discipline this whole project's own history argues for (a passing
check on the source is not proof the rendered artifact is correct).
"""
import re
import sys
from pathlib import Path

from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
DIST = ROOT / "dist"

FORBIDDEN = [
    r"meyklee",
    r"MEE2159",
    r"hkustmech",
    r"RTXG-KKGR",
    r"oauth/device",
    r"C:\\Users",
    r"user_code",
]

errors = []

for lang in ("EN", "ZH"):
    path = DIST / f"gstack-tutorial-4_{lang}.pptx"
    prs = Presentation(str(path))
    print(f"{lang}: {path.name} -- {len(prs.slides)} slides")

    all_text = []
    for si, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                for run in para.runs:
                    if run.text:
                        all_text.append(run.text)
    joined = "\n".join(all_text)

    for pattern in FORBIDDEN:
        m = re.search(pattern, joined, re.IGNORECASE)
        if m:
            errors.append(f"{lang}: forbidden pattern {pattern!r} found: ...{joined[max(0,m.start()-40):m.end()+40]}...")

    if len(prs.slides) != 20:
        errors.append(f"{lang}: expected 20 slides, got {len(prs.slides)}")

print()
if errors:
    print(f"FAIL -- {len(errors)} issue(s):")
    for e in errors:
        print(f"  - {e}")
    sys.exit(1)
print("PASS -- both decks 20/20 slides, zero forbidden identifiers in rendered text.")
